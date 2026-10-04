from io import BytesIO
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import UserCreationForm
from django.core.files.storage import default_storage
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .models import (
    AcademicCredential, Award, LanguageSkill, Publication, ResearchDirection,
    ResearchProject, ScientistProfile, TeachingSubject, WorkExperience,
)


def register_view(request):
    if request.user.is_authenticated:
        return redirect("profile_home")
    form = UserCreationForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        ScientistProfile.objects.create(user=user, full_name=user.username)
        login(request, user)
        return redirect("profile_home")
    return render(request, "registration/register.html", {"form": form})


def _profile(user):
    p, _ = ScientistProfile.objects.get_or_create(
        user=user, defaults={"full_name": user.get_full_name() or user.username, "email": user.email}
    )
    return p


@login_required
def profile_home(request):
    p = _profile(request.user)
    return render(request, "profiles/home.html", {
        "profile": p,
        "teaching_subjects": p.teaching_subjects.all(),
        "research_directions": p.research_directions.all(),
        "credentials": p.credentials.all().order_by("-year"),
        "languages": p.languages.all(),
        "work_experiences": p.work_experiences.all().order_by("-start_date"),
        "projects": p.projects.all().order_by("-start_year"),
        "publications": p.publications.all().order_by("-year", "title"),
        "awards": p.awards.all().order_by("-year"),
    })


def _date(value):
    from datetime import datetime
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _int(value):
    try:
        return int(value) if value else None
    except (ValueError, TypeError):
        return None


@login_required
@require_POST
def update_general(request):
    p = _profile(request.user)
    p.full_name = request.POST.get("person_name", "")
    p.gender = _int(request.POST.get("person_gender")) or 0
    p.dob = _date(request.POST.get("person_dob"))
    p.pob = request.POST.get("person_pob", "")
    p.hometown = request.POST.get("person_hometown", "")
    p.ethnicity = request.POST.get("person_ethnicity", "")
    p.position = request.POST.get("person_position", "")
    p.work_unit = request.POST.get("person_work_unit", "")
    p.address = request.POST.get("person_address", "")
    p.office_phone = request.POST.get("person_office_phone", "")
    p.home_phone = request.POST.get("person_home_phone", "")
    p.mobile_phone = request.POST.get("person_mobile_phone", "")
    p.fax = request.POST.get("person_fax", "")
    p.email = request.POST.get("person_email", "")
    p.citizen_id = request.POST.get("person_citizen_id", "")
    p.date_issue = _date(request.POST.get("person_date_issue"))
    p.place_issue = request.POST.get("person_place_issue", "")
    p.save()
    messages.success(request, "Đã lưu thông tin cơ bản.")
    return redirect("profile_home")


@login_required
@require_POST
def update_preferences(request):
    p = _profile(request.user)
    for field in ["accept_teaching", "accept_researching", "accept_committee", "accept_instructor"]:
        setattr(p, field, request.POST.get(field) == "on")
    p.save(update_fields=["accept_teaching", "accept_researching", "accept_committee", "accept_instructor"])
    messages.success(request, "Đã cập nhật mong muốn.")
    return redirect("profile_home")


@login_required
@require_POST
def add_teaching(request):
    TeachingSubject.objects.create(profile=_profile(request.user), major=request.POST.get("major", ""), subject=request.POST.get("subject", ""), description=request.POST.get("description", ""))
    return redirect("profile_home")


@login_required
@require_POST
def add_research_direction(request):
    ResearchDirection.objects.create(profile=_profile(request.user), major=request.POST.get("major", ""), direction=request.POST.get("direction", ""), description=request.POST.get("description", ""))
    return redirect("profile_home")


@login_required
@require_POST
def add_credential(request):
    AcademicCredential.objects.create(
        profile=_profile(request.user), kind=request.POST.get("kind", "degree"), level=request.POST.get("level", ""),
        major=request.POST.get("major", ""), institution=request.POST.get("institution", ""), year=_int(request.POST.get("year")),
        training_system=request.POST.get("training_system", ""), country=request.POST.get("country", ""), thesis=request.POST.get("thesis", ""),
    )
    return redirect("profile_home")


@login_required
@require_POST
def add_language(request):
    LanguageSkill.objects.create(profile=_profile(request.user), name=request.POST.get("name", ""), proficiency=request.POST.get("proficiency", ""), certificate=request.POST.get("certificate", ""))
    return redirect("profile_home")


@login_required
@require_POST
def add_work(request):
    current = request.POST.get("current") == "on"
    WorkExperience.objects.create(
        profile=_profile(request.user), start_date=_date(request.POST.get("start_date")), end_date=None if current else _date(request.POST.get("end_date")),
        current=current, institution=request.POST.get("institution", ""), position=request.POST.get("position", ""), description=request.POST.get("description", ""),
    )
    return redirect("profile_home")


@login_required
@require_POST
def add_project(request):
    raw = (request.POST.get("budget") or "").replace(".", "").replace(",", ".")
    try:
        budget = Decimal(raw) if raw else None
    except InvalidOperation:
        budget = None
    ResearchProject.objects.create(
        profile=_profile(request.user), title=request.POST.get("title", ""), start_year=_int(request.POST.get("start_year")), end_year=_int(request.POST.get("end_year")),
        level=request.POST.get("level", ""), ranking=request.POST.get("ranking", ""), position=request.POST.get("position", ""), budget=budget, budget_currency=request.POST.get("currency", "VND")
    )
    return redirect("profile_home")


@login_required
@require_POST
def add_publication(request):
    Publication.objects.create(
        profile=_profile(request.user), type=request.POST.get("type", "article"), title=request.POST.get("title", ""), location=request.POST.get("location", ""),
        year=_int(request.POST.get("year")), month=_int(request.POST.get("month")), journal=request.POST.get("journal", ""), doi=request.POST.get("doi", ""),
        url=request.POST.get("url", ""), description=request.POST.get("description", ""),
    )
    return redirect("profile_home")


@login_required
@require_POST
def add_award(request):
    Award.objects.create(profile=_profile(request.user), category=request.POST.get("category", "science_award"), name=request.POST.get("name", ""), organization=request.POST.get("organization", ""), year=_int(request.POST.get("year")), description=request.POST.get("description", ""))
    return redirect("profile_home")


@login_required
@require_POST
def delete_item(request, kind, pk):
    p = _profile(request.user)
    mapping = {
        "teaching": TeachingSubject, "research": ResearchDirection, "credential": AcademicCredential,
        "language": LanguageSkill, "work": WorkExperience, "project": ResearchProject, "publication": Publication, "award": Award,
    }
    model = mapping.get(kind)
    if model:
        get_object_or_404(model, pk=pk, profile=p).delete()
    return redirect("profile_home")


@login_required
@require_POST
def support_upload(request):
    f = request.FILES.get("support_file")
    if not f:
        messages.error(request, "Chưa chọn tệp.")
        return redirect("profile_home")
    suffix = Path(f.name).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        messages.error(request, "Chỉ hỗ trợ PDF hoặc DOCX.")
        return redirect("profile_home")
    saved = default_storage.save(f"support_uploads/{request.user.id}_{f.name}", f)
    messages.success(request, f"Đã tải lên {saved}. Bước trích xuất dữ liệu tự động sẽ được bổ sung ở phiên bản tiếp theo.")
    return redirect("profile_home")


def _set_doc_font(doc):
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Times New Roman"
            style.font.size = Pt(13)


@login_required
def export_cv(request):
    p = _profile(request.user)
    doc = Document()
    _set_doc_font(doc)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc")
    r.bold = True
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = title.add_run("LÝ LỊCH KHOA HỌC")
    rr.bold = True; rr.font.size = Pt(16)
    doc.add_paragraph("(Theo mẫu tại Thông tư số 08/2011/TT-BGDĐT ngày 17/02/2011 của Bộ trưởng Bộ GDĐT – Phụ lục V)").alignment = WD_ALIGN_PARAGRAPH.CENTER

    h = doc.add_paragraph(); run=h.add_run("I. LÝ LỊCH SƠ LƯỢC"); run.bold=True
    lines = [
        f"Họ và tên: {p.full_name}    Giới tính: {p.get_gender_display()}",
        f"Ngày, tháng, năm sinh: {p.dob or ''}    Nơi sinh: {p.pob}",
        f"Quê quán: {p.hometown}    Dân tộc: {p.ethnicity}",
        f"Chức vụ (hiện tại hoặc trước khi nghỉ hưu): {p.position}",
        f"Đơn vị công tác (hiện tại hoặc trước khi nghỉ hưu): {p.work_unit}",
        f"Chỗ ở riêng hoặc địa chỉ liên lạc: {p.address}",
        f"Điện thoại liên hệ: CQ: {p.office_phone}    NR: {p.home_phone}    DĐ: {p.mobile_phone}",
        f"Fax: {p.fax}    E-mail: {p.email}",
        f"Số CCCD: {p.citizen_id}    Ngày cấp: {p.date_issue or ''}    Nơi cấp: {p.place_issue}",
    ]
    for x in lines: doc.add_paragraph(x)

    h=doc.add_paragraph(); r=h.add_run("II. QUÁ TRÌNH ĐÀO TẠO"); r.bold=True
    for i,c in enumerate(p.credentials.all(),1):
        doc.add_paragraph(f"{i}. {c.get_kind_display()}: {c.level} - {c.major}; {c.institution}; {c.country}; Năm: {c.year or ''}; {c.thesis}")
    if p.languages.exists():
        doc.add_paragraph("Ngoại ngữ:")
        for i,l in enumerate(p.languages.all(),1): doc.add_paragraph(f"{i}. {l.name} - Mức độ sử dụng: {l.proficiency} - Chứng chỉ: {l.certificate}")

    doc.add_page_break()
    h=doc.add_paragraph(); r=h.add_run("III. QUÁ TRÌNH CÔNG TÁC CHUYÊN MÔN"); r.bold=True
    table=doc.add_table(rows=1, cols=3); table.style="Table Grid"
    for cell,txt in zip(table.rows[0].cells,["Thời gian","Nơi công tác","Công việc đảm nhiệm"]): cell.text=txt
    for w in p.work_experiences.all():
        cells=table.add_row().cells
        cells[0].text=f"{w.start_date} - {'Hiện tại' if w.current else (w.end_date or '')}"
        cells[1].text=w.institution
        cells[2].text=w.position

    h=doc.add_paragraph(); r=h.add_run("IV. QUÁ TRÌNH NGHIÊN CỨU KHOA HỌC"); r.bold=True
    doc.add_paragraph("1. Các đề tài nghiên cứu khoa học đã và đang tham gia:")
    t=doc.add_table(rows=1, cols=5); t.style="Table Grid"
    headers=["TT","Tên đề tài nghiên cứu","Năm bắt đầu/Năm hoàn thành","Đề tài cấp","Trách nhiệm tham gia trong đề tài"]
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    for i,proj in enumerate(p.projects.all(),1):
        cells=t.add_row().cells
        vals=[str(i),proj.title,f"{proj.start_year or ''}/{proj.end_year or ''}",proj.level,proj.position]
        for c,v in zip(cells,vals): c.text=v
    doc.add_paragraph("2. Các công trình khoa học đã công bố:")
    t=doc.add_table(rows=1, cols=4); t.style="Table Grid"
    for c,txt in zip(t.rows[0].cells,["TT","Tên công trình","Năm công bố","Tên tạp chí"]): c.text=txt
    for i,pub in enumerate(p.publications.all(),1):
        cells=t.add_row().cells
        vals=[str(i),pub.title,str(pub.year or ''),pub.journal or pub.location]
        for c,v in zip(cells,vals): c.text=v

    out=BytesIO(); doc.save(out); out.seek(0)
    resp=HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    resp["Content-Disposition"]='attachment; filename="Ly_lich_khoa_hoc.docx"'
    return resp
